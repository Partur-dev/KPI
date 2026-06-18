#!/usr/bin/env python3
"""Build the lab report in Markdown and DOCX formats."""

from __future__ import annotations

import json
import sqlite3
import textwrap
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


LAB_DIR = Path(__file__).resolve().parents[1]
ROOT_DIR = LAB_DIR.parent
OUTPUT_DIR = LAB_DIR / "output"
ASSET_DIR = LAB_DIR / "report_assets"
SUMMARY_PATH = OUTPUT_DIR / "etl_run_summary.json"
DB_PATH = OUTPUT_DIR / "olist_dw.sqlite"
REPORT_MD = LAB_DIR / "report.md"
REPORT_DOCX = LAB_DIR / "Olist_ETL_Data_Warehouse_Report.docx"
STAGE_DIAGRAM = ASSET_DIR / "stage_model.png"
WAREHOUSE_DIAGRAM = ASSET_DIR / "warehouse_model.png"


STAGE_DESCRIPTIONS = [
    ("stg_customers", "Покупці: ідентифікатори клієнтів, місто, штат, ZIP-префікс."),
    ("stg_geolocation", "Геолокація ZIP-префіксів: координати, місто та штат."),
    ("stg_orders", "Замовлення: статуси та часові мітки життєвого циклу."),
    ("stg_order_items", "Позиції замовлень: товар, продавець, ціна, доставка."),
    ("stg_order_payments", "Оплати: тип платежу, кількість платежів, сума."),
    ("stg_order_reviews", "Відгуки: оцінка, коментарі, дата створення і відповіді."),
    ("stg_products", "Товари: категорія, габарити, вага, довжини описів."),
    ("stg_sellers", "Продавці: ідентифікатор, місто, штат, ZIP-префікс."),
    ("stg_product_category_translation", "Переклад категорій товарів з португальської на англійську."),
]

WAREHOUSE_DESCRIPTIONS = [
    ("dim_date", "Календарний вимір для дат купівлі, доставки, відгуків та лімітів відправлення."),
    ("dim_location", "Вимір локацій за ZIP-префіксом, містом, штатом і середніми координатами."),
    ("dim_customer", "Вимір покупців із посиланням на локацію."),
    ("dim_seller", "Вимір продавців із посиланням на локацію."),
    ("dim_product", "Вимір товарів із фізичними характеристиками та категорією."),
    ("dim_category", "Вимір категорій з локальною та англомовною назвою."),
    ("dim_payment_type", "Вимір способів оплати."),
    ("dim_order_status", "Вимір статусів замовлення."),
    ("fact_order_items", "Факт продажів на рівні позиції замовлення: ціна, доставка, затримка."),
    ("fact_payments", "Факт оплат на рівні платежу: сума, тип, кількість частин."),
    ("fact_reviews", "Факт відгуків: оцінка, наявність коментаря, час відповіді."),
]

WAREHOUSE_DIMENSIONS = WAREHOUSE_DESCRIPTIONS[:8]
WAREHOUSE_FACTS = WAREHOUSE_DESCRIPTIONS[8:]


def load_font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    candidates = [
        "/System/Library/Fonts/Supplemental/Arial Bold.ttf" if bold else "/System/Library/Fonts/Supplemental/Arial.ttf",
        "/Library/Fonts/Arial Unicode.ttf",
        "/System/Library/Fonts/Supplemental/DejaVu Sans.ttf",
    ]
    for candidate in candidates:
        if candidate and Path(candidate).exists():
            return ImageFont.truetype(candidate, size=size)
    return ImageFont.load_default()


def rounded_box(draw: ImageDraw.ImageDraw, xy: tuple[int, int, int, int], fill: str, outline: str, text: str, font, text_fill="#111827") -> None:
    draw.rounded_rectangle(xy, radius=18, fill=fill, outline=outline, width=2)
    x1, y1, x2, y2 = xy
    lines = []
    for raw in text.split("\n"):
        wrapped = textwrap.wrap(raw, width=max(8, (x2 - x1) // 14))
        lines.extend(wrapped or [""])
    line_h = font.size + 6 if hasattr(font, "size") else 20
    total_h = line_h * len(lines)
    y = y1 + ((y2 - y1) - total_h) // 2
    for line in lines:
        bbox = draw.textbbox((0, 0), line, font=font)
        draw.text((x1 + ((x2 - x1) - (bbox[2] - bbox[0])) // 2, y), line, fill=text_fill, font=font)
        y += line_h


def arrow(draw: ImageDraw.ImageDraw, start: tuple[int, int], end: tuple[int, int], color: str = "#334155") -> None:
    draw.line([start, end], fill=color, width=3)
    ex, ey = end
    sx, sy = start
    if ex >= sx:
        points = [(ex, ey), (ex - 14, ey - 8), (ex - 14, ey + 8)]
    else:
        points = [(ex, ey), (ex + 14, ey - 8), (ex + 14, ey + 8)]
    draw.polygon(points, fill=color)


def create_stage_diagram() -> None:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    img = Image.new("RGB", (1800, 1040), "#FFFFFF")
    draw = ImageDraw.Draw(img)
    title_font = load_font(42, bold=True)
    box_font = load_font(24)
    small_font = load_font(20)
    draw.text((60, 40), "Модель Stage-зони Olist", fill="#0F172A", font=title_font)
    rounded_box(draw, (700, 150, 1100, 890), "#E8EEF5", "#2E74B5", "Stage zone\nSQLite\nraw/stg_* tables", title_font, "#0F172A")

    sources = [
        "customers.csv", "geolocation.csv", "orders.csv", "order_items.csv", "payments.csv",
        "reviews.csv", "products.csv", "sellers.csv", "category_translation.csv",
    ]
    left_positions = [(70, 150 + i * 95, 420, 220 + i * 95) for i in range(5)]
    right_positions = [(1380, 150 + i * 95, 1730, 220 + i * 95) for i in range(4)]
    for idx, xy in enumerate(left_positions):
        rounded_box(draw, xy, "#F8FAFC", "#CBD5E1", sources[idx], box_font)
        arrow(draw, (xy[2], (xy[1] + xy[3]) // 2), (700, (xy[1] + xy[3]) // 2))
    for offset, xy in enumerate(right_positions):
        idx = offset + 5
        rounded_box(draw, xy, "#F8FAFC", "#CBD5E1", sources[idx], box_font)
        arrow(draw, (xy[0], (xy[1] + xy[3]) // 2), (1100, (xy[1] + xy[3]) // 2))
    draw.text((620, 925), "Завантаження 1:1 з CSV, очищення BOM/порожніх значень, контроль кількості рядків", fill="#334155", font=small_font)
    img.save(STAGE_DIAGRAM)


def create_warehouse_diagram() -> None:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    img = Image.new("RGB", (2000, 1280), "#FFFFFF")
    draw = ImageDraw.Draw(img)
    title_font = load_font(42, bold=True)
    box_font = load_font(24)
    fact_font = load_font(26, bold=True)
    draw.text((60, 40), "Модель основного сховища: сузір'я фактів", fill="#0F172A", font=title_font)

    facts = {
        "fact_order_items\nsales, freight,\ndelivery delay": (760, 390, 1240, 540),
        "fact_payments\npayment value,\ninstallments": (760, 585, 1240, 735),
        "fact_reviews\nscore,\nanswer delay": (760, 780, 1240, 930),
    }
    dims = {
        "dim_date": (110, 170, 410, 260),
        "dim_customer": (110, 400, 410, 490),
        "dim_seller": (110, 650, 410, 740),
        "dim_location": (110, 900, 410, 990),
        "dim_product": (1590, 300, 1890, 390),
        "dim_category": (1590, 510, 1890, 600),
        "dim_payment_type": (1590, 720, 1890, 810),
        "dim_order_status": (1590, 930, 1890, 1020),
    }
    for text, xy in facts.items():
        rounded_box(draw, xy, "#E8EEF5", "#2E74B5", text, fact_font, "#0F172A")
    for text, xy in dims.items():
        rounded_box(draw, xy, "#F8FAFC", "#CBD5E1", text, box_font, "#111827")

    fact_centers = [(760, 465), (760, 660), (760, 855), (1240, 465), (1240, 660), (1240, 855)]
    for text, xy in dims.items():
        x1, y1, x2, y2 = xy
        center = ((x1 + x2) // 2, (y1 + y2) // 2)
        if x2 < 760:
            targets = [(760, 465), (760, 660), (760, 855)]
            for target in targets:
                arrow(draw, (x2, center[1]), target)
        else:
            targets = [(1240, 465), (1240, 660), (1240, 855)]
            for target in targets:
                arrow(draw, (x1, center[1]), target)
    img.save(WAREHOUSE_DIAGRAM)


def table_counts(conn: sqlite3.Connection, table_names: list[str]) -> dict[str, int]:
    return {name: conn.execute(f"SELECT COUNT(*) FROM {name}").fetchone()[0] for name in table_names}


def source_table() -> list[tuple[str, str, str]]:
    return [
        ("olist_customers_dataset.csv", "stg_customers", "Вимір покупців і локацій."),
        ("olist_geolocation_dataset.csv", "stg_geolocation", "Збагачення локацій координатами."),
        ("olist_orders_dataset.csv", "stg_orders", "Статуси та дати замовлень."),
        ("olist_order_items_dataset.csv", "stg_order_items", "Факт продажів на рівні позиції."),
        ("olist_order_payments_dataset.csv", "stg_order_payments", "Факт оплат."),
        ("olist_order_reviews_dataset.csv", "stg_order_reviews", "Факт відгуків."),
        ("olist_products_dataset.csv", "stg_products", "Вимір товарів."),
        ("olist_sellers_dataset.csv", "stg_sellers", "Вимір продавців."),
        ("product_category_name_translation.csv", "stg_product_category_translation", "Вимір категорій."),
    ]


def markdown_report(summary: dict, stage_counts: dict[str, int], warehouse_counts: dict[str, int]) -> str:
    lines = [
        "# Практикум №1-2: створення процедур завантаження даних",
        "",
        "**Предметна область:** електронна комерція Olist.",
        "**Виконавець:** студент/студентка ____________________, група __________.",
        "**Середовище:** Python 3, SQLite, CSV-файли з папки `data`.",
        "",
        "## Опис джерел даних",
        "",
        "| Файл | Stage-таблиця | Призначення | Рядків |",
        "|---|---:|---|---:|",
    ]
    for file_name, table, purpose in source_table():
        lines.append(f"| `{file_name}` | `{table}` | {purpose} | {stage_counts.get(table, 0):,} |")
    lines.extend([
        "",
        "## Stage-зона",
        "",
        "Stage-зона повторює структуру CSV-файлів і зберігає сирі дані перед трансформаціями.",
        "",
        "![Модель Stage-зони](report_assets/stage_model.png)",
        "",
        "| Таблиця | Опис | Рядків |",
        "|---|---|---:|",
    ])
    for table, description in STAGE_DESCRIPTIONS:
        lines.append(f"| `{table}` | {description} | {stage_counts.get(table, 0):,} |")
    lines.extend([
        "",
        "## Основне сховище",
        "",
        "Сховище побудоване як сузір'я фактів: спільні виміри використовуються трьома фактами.",
        "",
        "![Модель сховища](report_assets/warehouse_model.png)",
        "",
        "| Таблиця | Опис | Рядків |",
        "|---|---|---:|",
    ])
    for table, description in WAREHOUSE_DESCRIPTIONS:
        lines.append(f"| `{table}` | {description} | {warehouse_counts.get(table, 0):,} |")
    lines.extend([
        "",
        "## ETL-засоби",
        "",
        "- `scripts/etl_olist.py` виконує завантаження CSV у stage і трансформацію у warehouse.",
        "- `sql/01_stage_schema.sql` створює stage-зону.",
        "- `sql/02_warehouse_schema.sql` створює виміри, факти, ключі та індекси.",
        "- `sql/03_transform_load.sql` виконує очищення, нормалізацію, розрахунок метрик і `UPSERT`.",
        "- `sql/04_analysis_queries.sql` містить приклади аналітичних запитів.",
        "",
        "Передбачено повторне завантаження змінених і додаткових даних: stage перезавантажується з CSV, а warehouse оновлюється за природними ключами (`product_id`, `customer_id`, `seller_id`, `order_id + order_item_id`, `order_id + payment_sequential`, `review_id + order_id`).",
        "",
        "## Результати запуску",
        "",
        f"- Запуск ETL: `run_id={summary['run_id']}`, тривалість {summary['duration_seconds']} с.",
        "- Перевірки якості: відсутні факти без покупця, товару, продавця, типу платежу; оцінки відгуків у межах 1-5.",
        "",
        "| Перевірка | Результат |",
        "|---|---:|",
    ])
    for check, value in summary["quality_checks"].items():
        lines.append(f"| `{check}` | {value} |")
    lines.extend([
        "",
        "### Топ категорій за виручкою",
        "",
        "| Категорія | Позицій | Виручка | Середня оцінка |",
        "|---|---:|---:|---:|",
    ])
    for row in summary["top_category_metrics"]:
        lines.append(
            f"| {row['category']} | {row['order_items']:,} | {row['total_revenue']:,.2f} | {row['avg_review_score']:.2f} |"
        )
    lines.extend([
        "",
        "## Відповіді на питання самоперевірки",
        "",
        "**ETL** - процес Extract, Transform, Load: отримання даних із джерел, приведення до потрібної якості та структури, завантаження у цільову БД.",
        "",
        "**Процес ETL базується** на аналізі джерел, правилах очищення, бізнес-ключах, перевірках якості, трансформаціях і контрольованому завантаженні у модель сховища.",
        "",
        "**Зірка** проста для розуміння і швидка для аналітичних запитів, але може дублювати атрибути. **Сніжинка** зменшує дублювання через нормалізацію вимірів, але збільшує кількість JOIN і складність моделі.",
        "",
        "**Підходи до багатовимірного моделювання:** Kimball bottom-up із вітринами і конформними вимірами; Inmon top-down із корпоративним сховищем; Data Vault для історизації та інтеграції; OLAP-куби/табулярні моделі для аналітичного шару.",
        "",
        "## Висновок",
        "",
        "У роботі спроектовано stage-зону та основне сховище для Olist, створено ETL-скрипти, виконано повне завантаження даних і перевірено ідемпотентне повторне завантаження. Сховище містить понад п'ять вимірів і три таблиці фактів, що дозволяє аналізувати продажі, оплати, відгуки, категорії, географію та доставку.",
        "",
    ])
    return "\n".join(lines)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text: str, bold: bool = False) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(str(text))
    run.bold = bold
    run.font.name = "Calibri"
    run.font.size = Pt(9)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def add_table(doc: Document, headers: list[str], rows: list[list[object]], widths: list[float]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_text(cell, header, bold=True)
        set_cell_shading(cell, "F2F4F7")
        cell.width = Inches(widths[idx])
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            set_cell_text(cells[idx], value)
            cells[idx].width = Inches(widths[idx])
    doc.add_paragraph()


def add_bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(4)
    p.add_run(text)


def add_numbered(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(4)
    p.add_run(text)


def build_docx(summary: dict, stage_counts: dict[str, int], warehouse_counts: dict[str, int]) -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(11)
    styles["Normal"].paragraph_format.space_after = Pt(6)
    styles["Normal"].paragraph_format.line_spacing = 1.10
    for style_name, size in [("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 12)]:
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(46, 116, 181 if style_name != "Heading 3" else 120)
        style.paragraph_format.space_before = Pt(16 if style_name == "Heading 1" else 12)
        style.paragraph_format.space_after = Pt(8 if style_name == "Heading 1" else 6)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Практикум №1-2\nСтворення процедур завантаження даних")
    run.bold = True
    run.font.size = Pt(20)
    run.font.name = "Calibri"
    run.font.color.rgb = RGBColor(15, 23, 42)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("Курс: Аналіз даних в інформаційних системах\nДатасет: Olist Brazilian E-Commerce").italic = True
    doc.add_paragraph()
    add_table(
        doc,
        ["Поле", "Значення"],
        [
            ["Виконавець", "студент/студентка ____________________"],
            ["Група", "__________"],
            ["Інструменти", "Python 3, SQLite, CSV, python-docx"],
            ["Результат", "Stage-зона, dimensional warehouse, ETL-скрипти, звіт"],
        ],
        [1.8, 4.4],
    )

    doc.add_heading("1. Опис джерел даних", level=1)
    doc.add_paragraph(
        "Для лабораторних робіт використано класичний набір відкритих даних Olist. "
        "Джерела мають спільні виміри: замовлення, товари, покупці, продавці, категорії, дати та локації."
    )
    add_table(
        doc,
        ["Файл", "Stage-таблиця", "Призначення", "Рядків"],
        [[file_name, table, purpose, f"{stage_counts.get(table, 0):,}"] for file_name, table, purpose in source_table()],
        [2.1, 1.5, 2.5, 0.8],
    )

    doc.add_heading("2. Модель Stage-зони", level=1)
    doc.add_paragraph(
        "Stage-зона зберігає сирі дані 1:1 із CSV. На цьому етапі виконується технічне очищення: "
        "зняття BOM, перетворення порожніх рядків на NULL, фіксація кількості рядків і часу завантаження."
    )
    doc.add_picture(str(STAGE_DIAGRAM), width=Inches(6.3))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_table(
        doc,
        ["Таблиця", "Опис", "Рядків"],
        [[table, description, f"{stage_counts.get(table, 0):,}"] for table, description in STAGE_DESCRIPTIONS],
        [1.8, 3.9, 0.9],
    )

    doc.add_heading("3. Модель основного сховища", level=1)
    doc.add_paragraph(
        "Основне сховище спроектовано як сузір'я фактів. Три фактові таблиці використовують конформні виміри "
        "для аналізу продажів, оплат, відгуків, доставки, географії та товарних категорій."
    )
    doc.add_picture(str(WAREHOUSE_DIAGRAM), width=Inches(6.4))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_heading("3.1. Виміри", level=2)
    add_table(
        doc,
        ["Таблиця", "Опис", "Рядків"],
        [[table, description, f"{warehouse_counts.get(table, 0):,}"] for table, description in WAREHOUSE_DIMENSIONS],
        [1.7, 4.0, 0.9],
    )
    doc.add_heading("3.2. Таблиці фактів", level=2)
    add_table(
        doc,
        ["Таблиця", "Опис", "Рядків"],
        [[table, description, f"{warehouse_counts.get(table, 0):,}"] for table, description in WAREHOUSE_FACTS],
        [1.7, 4.0, 0.9],
    )

    doc.add_heading("4. ETL-засоби", level=1)
    add_bullet(doc, "scripts/etl_olist.py - керує повним процесом Extract, Transform, Load.")
    add_bullet(doc, "sql/01_stage_schema.sql - створює stage-таблиці.")
    add_bullet(doc, "sql/02_warehouse_schema.sql - створює виміри, факти, ключі та індекси.")
    add_bullet(doc, "sql/03_transform_load.sql - виконує нормалізацію, розрахунки і UPSERT.")
    add_bullet(doc, "sql/04_analysis_queries.sql - містить приклади аналітичних запитів.")
    doc.add_paragraph(
        "Можливість завантаження змінених і додаткових даних реалізована через природні ключі та "
        "операції ON CONFLICT DO UPDATE/DO NOTHING. Повторний запуск ETL не дублює записи, а оновлює "
        "атрибути вимірів і фактів, якщо відповідні CSV змінилися."
    )
    add_numbered(doc, "Extract: читання дев'яти CSV-файлів у stage-таблиці.")
    add_numbered(doc, "Transform: очищення тексту, нормалізація статусів, категорій, локацій і дат.")
    add_numbered(doc, "Load: наповнення вимірів, потім фактів із зовнішніми ключами.")
    add_numbered(doc, "Validate: контроль кількості рядків і ключових помилок якості.")

    doc.add_heading("5. Результати завантаження", level=1)
    doc.add_paragraph(
        f"Останній запуск ETL: run_id={summary['run_id']}, тривалість {summary['duration_seconds']} с. "
        "Повторний запуск без --reset підтвердив ідемпотентність: кількості рядків у warehouse не збільшилися."
    )
    add_table(
        doc,
        ["Перевірка", "Результат"],
        [[key, value] for key, value in summary["quality_checks"].items()],
        [4.8, 1.2],
    )
    doc.add_heading("5.1. Топ категорій за виручкою", level=2)
    add_table(
        doc,
        ["Категорія", "Позицій", "Виручка", "Середня оцінка"],
        [
            [
                row["category"],
                f"{row['order_items']:,}",
                f"{row['total_revenue']:,.2f}",
                f"{row['avg_review_score']:.2f}",
            ]
            for row in summary["top_category_metrics"]
        ],
        [2.3, 1.0, 1.4, 1.3],
    )

    doc.add_heading("6. Відповіді на питання самоперевірки", level=1)
    doc.add_paragraph("ETL - це процес Extract, Transform, Load: отримання даних із джерел, їх перетворення та завантаження у цільове сховище.")
    doc.add_paragraph("Процес ETL базується на аналізі джерел, бізнес-правилах, ключах інтеграції, перевірках якості, логуванні та контрольованому завантаженні.")
    doc.add_paragraph("Схема зірка проста для аналітики і швидка в запитах, але допускає дублювання атрибутів. Сніжинка зменшує дублювання завдяки нормалізації, але ускладнює JOIN і підтримку.")
    doc.add_paragraph("Підходи до багатовимірного моделювання: Kimball bottom-up, Inmon top-down, Data Vault, OLAP-куби та табулярні моделі семантичного шару.")

    doc.add_heading("7. Висновок", level=1)
    doc.add_paragraph(
        "У роботі спроектовано stage-зону та основне сховище для Olist, створено ETL-скрипти, "
        "виконано повне завантаження даних і перевірено повторне ідемпотентне завантаження. "
        "Сховище містить понад п'ять вимірів і три фактові таблиці, що достатньо для подальшої курсової роботи."
    )

    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Olist ETL Data Warehouse Lab")

    doc.save(REPORT_DOCX)


def main() -> None:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    summary = json.loads(SUMMARY_PATH.read_text(encoding="utf-8"))
    stage_tables = [table for table, _ in STAGE_DESCRIPTIONS]
    warehouse_tables = [table for table, _ in WAREHOUSE_DESCRIPTIONS]
    with sqlite3.connect(DB_PATH) as conn:
        stage_counts = table_counts(conn, stage_tables)
        warehouse_counts = table_counts(conn, warehouse_tables)
    create_stage_diagram()
    create_warehouse_diagram()
    REPORT_MD.write_text(markdown_report(summary, stage_counts, warehouse_counts), encoding="utf-8")
    build_docx(summary, stage_counts, warehouse_counts)
    print(json.dumps({
        "markdown": str(REPORT_MD),
        "docx": str(REPORT_DOCX),
        "stage_diagram": str(STAGE_DIAGRAM),
        "warehouse_diagram": str(WAREHOUSE_DIAGRAM),
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
