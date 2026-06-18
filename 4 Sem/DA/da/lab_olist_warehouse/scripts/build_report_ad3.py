#!/usr/bin/env python3
"""Build the Data Studio dashboard lab report in Markdown and DOCX formats."""

from __future__ import annotations

import json
import shutil
import sqlite3
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


LAB_DIR = Path(__file__).resolve().parents[1]
OUTPUT_DIR = LAB_DIR / "output"
ASSET_DIR = LAB_DIR / "report_assets"
DB_PATH = OUTPUT_DIR / "olist_dw.sqlite"
REPORT_MD = LAB_DIR / "report_ad3.md"
REPORT_DOCX = LAB_DIR / "Olist_Data_Studio_Dashboard_Report.docx"
DASHBOARD_TOP = ASSET_DIR / "dashboard_overview_top.png"
DASHBOARD_CHARTS = ASSET_DIR / "dashboard_overview_charts.png"


DATA_SOURCE_FIELDS = [
    ("purchase_date", "Вимір", "Date", "Фільтр періоду та часові графіки"),
    ("year", "Вимір", "Number", "Групування за роком"),
    ("year_month", "Вимір", "Text / Year Month", "Колонки pivot table і місячні підписи"),
    ("month", "Вимір", "Number", "Сортування місяців"),
    ("category", "Вимір", "Text", "Фільтр категорій, рядки pivot table, bar chart"),
    ("customer_state", "Вимір", "Text", "Фільтр штату покупця, рядки pivot table"),
    ("customer_city", "Вимір", "Text", "Деталізація локації покупця"),
    ("seller_state", "Вимір", "Text", "Географія продавця"),
    ("order_status", "Вимір", "Text", "Фільтр статусу замовлення"),
    ("order_id", "Вимір / база метрики", "Text", "COUNT_DISTINCT(order_id)"),
    ("order_item_id", "Вимір / база метрики", "Number", "COUNT(order_item_id)"),
    ("price", "Метрика", "Currency / Number", "SUM(price) або варіант параметра"),
    ("freight_value", "Метрика", "Currency / Number", "SUM(freight_value) або варіант параметра"),
    ("revenue", "Метрика", "Currency / Number", "SUM(revenue)"),
    ("days_to_delivery", "Метрика", "Number", "AVG(days_to_delivery)"),
    ("delivery_delay_days", "Метрика", "Number", "AVG(delivery_delay_days)"),
    ("was_delivered_late", "Метрика", "Percent base", "AVG(was_delivered_late)"),
    ("delivery_status", "Вимір", "Text", "Сектори кругової діаграми"),
    ("review_score", "Метрика", "Number", "AVG(review_score)"),
]


VISUALIZATIONS = [
    (
        "KPI cards",
        "Без виміру",
        "SUM(revenue), COUNT_DISTINCT(order_id), AVG(days_to_delivery), AVG(was_delivered_late)",
        "Показують загальну виручку, кількість унікальних замовлень, середній час доставки та частку запізнень.",
    ),
    (
        "Pivot table",
        "Rows: category, customer_state; Columns: year_month",
        "SUM(revenue)",
        "Виконує вимогу до табличного звіту з трьома вимірами, включно з часом.",
    ),
    (
        "Pie chart",
        "delivery_status",
        "SUM(revenue)",
        "Показує частку виручки за своєчасними, запізнілими та недоставленими замовленнями.",
    ),
    (
        "Line / area chart",
        "year_month",
        "SUM(selected_metric_value)",
        "Показує місячну динаміку та керується параметром звіту p_metric.",
    ),
    (
        "Bar chart",
        "category",
        "SUM(revenue)",
        "Показує провідні товарні категорії за виручкою.",
    ),
]


FILTERS_AND_PARAMETERS = [
    ("p_metric", "Parameter control", "Перемикає місячний графік між метриками revenue, price і freight."),
    ("category", "Drop-down filter", "Фільтрує всі компоненти dashboard за товарною категорією."),
    ("customer_state", "Drop-down filter", "Фільтрує всі компоненти dashboard за штатом покупця."),
    ("purchase_date", "Date range control", "Обмежує всі візуалізації вибраним періодом."),
]


def find_cleanshot(marker: str) -> Path:
    media_dir = Path.home() / "Library/Application Support/CleanShot/media"
    matches = sorted(media_dir.rglob(f"CleanShot 2026-06-15 at 9*37*{marker}@2x.png"))
    if not matches:
        raise FileNotFoundError(f"Could not find dashboard screenshot for marker {marker}")
    return matches[0]


def copy_dashboard_assets() -> None:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    shutil.copy2(find_cleanshot("51"), DASHBOARD_TOP)
    shutil.copy2(find_cleanshot("56"), DASHBOARD_CHARTS)


def query_one(conn: sqlite3.Connection, sql: str) -> tuple:
    return conn.execute(sql).fetchone()


def query_all(conn: sqlite3.Connection, sql: str) -> list[tuple]:
    return conn.execute(sql).fetchall()


def load_dashboard_metrics() -> dict:
    db_uri = f"file:{DB_PATH}?mode=ro&immutable=1"
    with sqlite3.connect(db_uri, uri=True) as conn:
        overview = query_one(
            conn,
            """
            SELECT
                MIN(d.full_date),
                MAX(d.full_date),
                COUNT(*),
                COUNT(DISTINCT f.order_id),
                ROUND(SUM(f.total_item_value), 2),
                ROUND(AVG(f.days_to_delivery), 2),
                ROUND(100.0 * AVG(f.was_delivered_late), 2)
            FROM fact_order_items f
            JOIN dim_date d ON d.date_key = f.purchase_date_key
            """,
        )
        top_categories = query_all(
            conn,
            """
            SELECT
                c.category_name_english,
                COUNT(*) AS items,
                ROUND(SUM(f.total_item_value), 2) AS revenue
            FROM fact_order_items f
            JOIN dim_product p ON p.product_key = f.product_key
            JOIN dim_category c ON c.category_key = p.category_key
            GROUP BY c.category_name_english
            ORDER BY revenue DESC
            LIMIT 5
            """,
        )
        delivery_status = query_all(
            conn,
            """
            SELECT
                CASE
                    WHEN f.was_delivered_late = 1 THEN 'Late'
                    WHEN f.was_delivered_late = 0 THEN 'On time'
                    ELSE 'Not delivered'
                END AS delivery_status,
                ROUND(SUM(f.total_item_value), 2) AS revenue
            FROM fact_order_items f
            GROUP BY delivery_status
            ORDER BY revenue DESC
            """,
        )
    return {
        "date_min": overview[0],
        "date_max": overview[1],
        "items": overview[2],
        "orders": overview[3],
        "revenue": overview[4],
        "avg_delivery_days": overview[5],
        "late_pct": overview[6],
        "top_categories": top_categories,
        "delivery_status": delivery_status,
    }


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text: object, bold: bool = False, size: int = 9) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(str(text))
    run.bold = bold
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def add_table(doc: Document, headers: list[str], rows: list[list[object]], widths: list[float]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_text(cell, header, bold=True)
        set_cell_shading(cell, "F2F4F7")
        cell.width = Inches(widths[idx])
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            set_cell_text(cells[idx], value)
            cells[idx].width = Inches(widths[idx])
    doc.add_paragraph()


def add_bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(4)
    p.add_run(text)


def add_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text)
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(75, 85, 99)


def add_picture(doc: Document, path: Path, width: float, caption: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width))
    add_caption(doc, caption)


def configure_styles(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(11)
    styles["Normal"].paragraph_format.space_after = Pt(6)
    styles["Normal"].paragraph_format.line_spacing = 1.10
    for style_name, size in [("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 12)]:
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(46, 116, 181 if style_name != "Heading 3" else 120)
        style.paragraph_format.space_before = Pt(16 if style_name == "Heading 1" else 12)
        style.paragraph_format.space_after = Pt(8 if style_name == "Heading 1" else 6)


def markdown_report(metrics: dict) -> str:
    lines = [
        "# Практикум №3: побудова звіту в Data Studio",
        "",
        "**Предметна область:** електронна комерція Olist.",
        "**Виконавець:** студент/студентка ____________________, група __________.",
        "**BI-інструмент:** Google Data Studio / Looker Studio.",
        "**Джерело:** експортована аналітична вітрина з SQLite-сховища, створеного в АД1-АD2.",
        "",
        "## DataSource",
        "",
        "DataSource побудовано на плоскій аналітичній таблиці `sales_dashboard`, яка об'єднує факт продажів з вимірами дати, категорії, покупця, продавця, статусу замовлення та відгуку.",
        "",
        "| Поле | Роль | Тип | Використання |",
        "|---|---|---|---|",
    ]
    for row in DATA_SOURCE_FIELDS:
        lines.append("| " + " | ".join(f"`{value}`" if idx == 0 else value for idx, value in enumerate(row)) + " |")
    lines.extend(
        [
            "",
            "## Dashboard",
            "",
            f"Період даних: {metrics['date_min']} - {metrics['date_max']}. Позицій замовлень: {metrics['items']:,}; унікальних замовлень: {metrics['orders']:,}; виручка: {metrics['revenue']:,.2f}.",
            "",
            "![Dashboard controls and pivot](report_assets/dashboard_overview_top.png)",
            "",
            "![Dashboard charts](report_assets/dashboard_overview_charts.png)",
            "",
            "## Опис візуалізацій",
            "",
            "| Візуалізація | Виміри | Метрики | Призначення |",
            "|---|---|---|---|",
        ]
    )
    for row in VISUALIZATIONS:
        lines.append("| " + " | ".join(row) + " |")
    lines.extend(
        [
            "",
            "## Фільтри та параметри",
            "",
            "| Елемент | Тип | Призначення |",
            "|---|---|---|",
        ]
    )
    for row in FILTERS_AND_PARAMETERS:
        lines.append("| " + " | ".join(f"`{value}`" if idx == 0 else value for idx, value in enumerate(row)) + " |")
    lines.extend(
        [
            "",
            "## Відповіді на питання самоперевірки",
            "",
            "**Структура аналітичних запитів:** SELECT вимірів і агрегованих метрик, FROM фактова таблиця, JOIN вимірів, WHERE фільтри, GROUP BY виміри, HAVING умови по агрегатах, ORDER BY сортування, LIMIT для top-N.",
            "",
            "**Переваги аналітичних запитів:** швидке агрегування великих даних, аналіз у розрізі часу, категорій і географії, побудова трендів, top-N, KPI та виявлення проблемних сегментів.",
            "",
        ]
    )
    return "\n".join(lines)


def build_docx(metrics: dict) -> None:
    doc = Document()
    configure_styles(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Практикум №3\nВізуалізація даних у Data Studio")
    run.bold = True
    run.font.size = Pt(20)
    run.font.name = "Calibri"
    run.font.color.rgb = RGBColor(15, 23, 42)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("Курс: Аналіз даних в інформаційних системах\nДатасет: Olist Brazilian E-Commerce").italic = True
    doc.add_paragraph()
    add_table(
        doc,
        ["Поле", "Значення"],
        [
            ["Виконавець", "студент/студентка ____________________"],
            ["Група", "__________"],
            ["BI-інструмент", "Google Data Studio / Looker Studio"],
            ["Джерело даних", "Аналітична CSV/Sheets-вітрина з SQLite-сховища АД1-АD2"],
            ["Результат", "DataSource, dashboard, фільтри, параметр, табличний і графічні звіти"],
        ],
        [1.8, 4.6],
    )

    doc.add_heading("1. DataSource", level=1)
    doc.add_paragraph(
        "Для побудови звіту використано аналітичну вітрину продажів Olist, експортовану зі сховища "
        "SQLite, створеного в АД1-АD2. Вітрина об'єднує `fact_order_items` з вимірами дати, категорії, "
        "локації покупця, локації продавця, статусу замовлення та відгуку."
    )
    add_table(
        doc,
        ["Показник", "Значення"],
        [
            ["Період даних", f"{metrics['date_min']} - {metrics['date_max']}"],
            ["Позицій замовлень", f"{metrics['items']:,}"],
            ["Унікальних замовлень", f"{metrics['orders']:,}"],
            ["Загальна виручка", f"{metrics['revenue']:,.2f}"],
            ["Середній час доставки", f"{metrics['avg_delivery_days']:.2f} днів"],
            ["Частка запізнілих доставок", f"{metrics['late_pct']:.2f}%"],
        ],
        [2.3, 3.9],
    )
    doc.add_heading("1.1. Поля DataSource", level=2)
    add_table(
        doc,
        ["Поле", "Роль", "Тип", "Використання"],
        [list(row) for row in DATA_SOURCE_FIELDS],
        [1.5, 1.25, 1.3, 2.45],
    )
    doc.add_heading("1.2. Розрахункові метрики", level=2)
    add_table(
        doc,
        ["Назва", "Формула / агрегація", "Призначення"],
        [
            ["Total Revenue", "SUM(revenue)", "Основна фінансова метрика дашборду."],
            ["Orders Count", "COUNT_DISTINCT(order_id)", "Кількість унікальних замовлень."],
            ["Avg Delivery Days", "AVG(days_to_delivery)", "Середня тривалість доставки."],
            ["Late Delivery %", "AVG(was_delivered_late)", "Частка доставок після планової дати."],
            ["selected_metric_value", "CASE by p_metric", "Параметрична метрика для лінійного графіка."],
        ],
        [1.55, 1.95, 2.7],
    )

    doc.add_heading("2. Екранні форми Dashboard", level=1)
    doc.add_paragraph(
        "Dashboard містить загальні KPI, фільтри, параметр вибору метрики, pivot table, кругову діаграму, "
        "лінійний графік за місяцями та стовпчиковий звіт за категоріями."
    )
    add_picture(doc, DASHBOARD_TOP, 6.4, "Рисунок 1 - верхня частина dashboard: фільтри, KPI та pivot table.")
    add_picture(doc, DASHBOARD_CHARTS, 6.4, "Рисунок 2 - графічні звіти dashboard: pie, monthly line chart та category bar chart.")

    doc.add_heading("3. Опис звітів", level=1)
    add_table(
        doc,
        ["Візуалізація", "Виміри", "Метрики", "Призначення"],
        [list(row) for row in VISUALIZATIONS],
        [1.35, 1.75, 1.85, 1.55],
    )
    doc.add_paragraph(
        "Кругова діаграма показує, що основна частина виручки припадає на замовлення зі статусом доставки "
        "`On time`. Лінійний графік демонструє місячну динаміку обраної параметром метрики. Стовпчиковий "
        "графік показує найсильніші товарні категорії; у топі знаходяться `health_beauty`, `watches_gifts`, "
        "`bed_bath_table`, `sports_leisure` та `computers_accessories`."
    )
    doc.add_heading("3.1. Top 5 категорій за виручкою", level=2)
    add_table(
        doc,
        ["Категорія", "Позицій", "Виручка"],
        [[category, f"{items:,}", f"{revenue:,.2f}"] for category, items, revenue in metrics["top_categories"]],
        [2.6, 1.2, 1.6],
    )
    doc.add_heading("3.2. Виручка за статусом доставки", level=2)
    add_table(
        doc,
        ["Статус доставки", "Виручка"],
        [[status, f"{revenue:,.2f}"] for status, revenue in metrics["delivery_status"]],
        [2.6, 1.8],
    )

    doc.add_heading("4. Фільтри, зв'язаність і параметри", level=1)
    add_table(
        doc,
        ["Елемент", "Тип", "Призначення"],
        [list(row) for row in FILTERS_AND_PARAMETERS],
        [1.7, 1.7, 3.0],
    )
    doc.add_paragraph(
        "Фільтри застосовуються до всіх візуалізацій dashboard, тому таблиця, KPI та графіки працюють як "
        "зв'язаний інтерактивний звіт. Параметр `p_metric` керує розрахунковим полем `selected_metric_value`, "
        "що дозволяє без створення окремих графіків перемикати аналіз між виручкою, ціною та доставкою."
    )
    doc.add_paragraph(
        "Географічна карта не додавалась, оскільки для коректної роботи Data Studio потрібне додаткове "
        "перетворення бразильських штатів у географічний тип. Географічний вимір усе одно використано через "
        "фільтр `customer_state` та pivot table у розрізі штатів покупців."
    )

    doc.add_heading("5. Відповіді на питання самоперевірки", level=1)
    doc.add_paragraph(
        "Структура аналітичного запиту: SELECT вимірів і агрегованих метрик, FROM фактова таблиця, JOIN "
        "вимірів, WHERE фільтри, GROUP BY виміри, HAVING умови по агрегатах, ORDER BY сортування та LIMIT "
        "для top-N звітів."
    )
    doc.add_paragraph(
        "Аналітичні запити дають змогу швидко агрегувати великі обсяги даних, порівнювати показники у розрізі "
        "часу, категорій і географії, будувати тренди, top-N звіти, KPI та знаходити проблемні сегменти, "
        "наприклад запізнілі доставки або категорії з низькими оцінками."
    )

    doc.add_heading("6. Висновок", level=1)
    doc.add_paragraph(
        "У роботі створено DataSource до сховища Olist, побудовано dashboard у Data Studio та налаштовано "
        "обов'язкові елементи: табличний звіт з трьома вимірами, кругову діаграму, лінійний графік по місяцях, "
        "top-звіт за категоріями, фільтри та параметричну метрику. Dashboard описує джерело даних через продажі, "
        "часову динаміку, категорії товарів, статус доставки та географію покупців."
    )

    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Olist Data Studio Dashboard Lab")

    doc.save(REPORT_DOCX)


def main() -> None:
    copy_dashboard_assets()
    metrics = load_dashboard_metrics()
    REPORT_MD.write_text(markdown_report(metrics), encoding="utf-8")
    build_docx(metrics)
    print(
        json.dumps(
            {
                "markdown": str(REPORT_MD),
                "docx": str(REPORT_DOCX),
                "dashboard_top": str(DASHBOARD_TOP),
                "dashboard_charts": str(DASHBOARD_CHARTS),
            },
            ensure_ascii=False,
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
